import streamlit as st
import json
import subprocess
import tempfile
import os
import re
from datetime import datetime

st.set_page_config(
    page_title="開発供給実施計画 申請書作成アシスタント",
    page_icon="📋",
    layout="wide",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Noto+Sans+JP:wght@300;400;500;700&display=swap');
html, body, [class*="css"] { font-family: 'Noto Sans JP', sans-serif; }
.main-header {
    background: linear-gradient(135deg, #1a237e 0%, #283593 50%, #1565c0 100%);
    color: white; padding: 2rem 2.5rem; border-radius: 12px;
    margin-bottom: 2rem; box-shadow: 0 4px 20px rgba(26,35,126,0.3);
}
.main-header h1 { font-size: 1.5rem; font-weight: 700; margin: 0; }
.main-header p  { font-size: 0.88rem; opacity: 0.85; margin: 0.4rem 0 0; }
.progress-bar-wrap { background: #e8eaf6; border-radius: 8px; height: 8px; margin: 1rem 0 0.4rem; }
.progress-bar-fill { background: linear-gradient(90deg, #42a5f5, #1e88e5); border-radius: 8px; height: 8px; }
.chat-user { background: #e3f2fd; border-left: 4px solid #1e88e5; padding: 0.8rem 1rem; border-radius: 0 8px 8px 0; margin: 0.6rem 0; white-space: pre-wrap; }
.chat-ai   { background: #f5f5f5; border-left: 4px solid #78909c; padding: 0.8rem 1rem; border-radius: 0 8px 8px 0; margin: 0.6rem 0; white-space: pre-wrap; }
.section-preview { background: white; border: 1px solid #e0e0e0; border-radius: 8px; padding: 0.8rem 1rem; margin: 0.4rem 0; font-size: 0.85rem; }
.section-title { font-weight: 700; color: #1a237e; font-size: 0.82rem; margin-bottom: 0.3rem; }
.cbadge { background:#e8f5e9;border:1px solid #a5d6a7;color:#2e7d32;padding:0.15rem 0.6rem;border-radius:20px;font-size:0.72rem;font-weight:600; }
.ibadge { background:#fff3e0;border:1px solid #ffcc80;color:#e65100;padding:0.15rem 0.6rem;border-radius:20px;font-size:0.72rem;font-weight:600; }
.provider-badge { display:inline-block; padding:0.2rem 0.8rem; border-radius:20px; font-size:0.78rem; font-weight:600; margin-left:0.5rem; }
.badge-claude  { background:#e8eaf6; color:#3949ab; border:1px solid #9fa8da; }
.badge-gemini  { background:#e8f5e9; color:#2e7d32; border:1px solid #a5d6a7; }
.badge-openai  { background:#e0f2f1; color:#00695c; border:1px solid #80cbc4; }
.badge-openrouter { background:#fce4ec; color:#c62828; border:1px solid #ef9a9a; }
.badge-groq      { background:#fff8e1; color:#e65100; border:1px solid #ffcc02; }
</style>
""", unsafe_allow_html=True)

# ─── プロバイダー設定 ─────────────────────────────────────────────────────────
PROVIDERS = {
    "Claude (Anthropic)": {
        "id": "claude",
        "badge": "badge-claude",
        "models": [
            "claude-sonnet-4-20250514",
            "claude-opus-4-20250514",
            "claude-haiku-4-5-20251001",
        ],
        "key_placeholder": "sk-ant-...",
        "key_url": "https://console.anthropic.com",
    },
    "Gemini (Google)": {
        "id": "gemini",
        "badge": "badge-gemini",
        "models": [
            "gemini-2.0-flash",
            "gemini-2.5-pro-preview-05-06",
            "gemini-1.5-pro",
        ],
        "key_placeholder": "AIza...",
        "key_url": "https://aistudio.google.com/app/apikey",
    },
    "ChatGPT (OpenAI)": {
        "id": "openai",
        "badge": "badge-openai",
        "models": [
            "gpt-4o",
            "gpt-4o-mini",
            "gpt-4-turbo",
            "o3-mini",
        ],
        "key_placeholder": "sk-...",
        "key_url": "https://platform.openai.com/api-keys",
    },
    "Groq": {
        "id": "groq",
        "badge": "badge-groq",
        "models": [
            # ── 高速・低コスト（おすすめ） ──────────────────────
            "llama-3.3-70b-versatile",        # 汎用・高品質  $0.59/$0.79
            "llama-3.1-8b-instant",            # 超高速・格安  $0.05/$0.08
            "qwen/qwen3-32b",                  # 推論強化      $0.29/$0.59
            # ── GPT-OSS（OpenAI公式OSS） ─────────────────────
            "openai/gpt-oss-120b",             # 高性能        $0.15/$0.60
            "openai/gpt-oss-20b",              # バランス型    $0.075/$0.30
            # ── Meta Llama 4 ─────────────────────────────────
            "meta-llama/llama-4-scout-17b-16e-instruct",  # Llama4 Scout $0.11/$0.34
            # ── Kimi K2（長文対応） ───────────────────────────
            "moonshotai/kimi-k2-instruct-0905",# 長文・高品質  $1.00/$3.00
            # ── 推論モデル ────────────────────────────────────
            "qwen-qwq-32b",                    # 推論特化
            "deepseek-r1-distill-llama-70b",   # 推論特化
        ],
        "key_placeholder": "gsk_...",
        "key_url": "https://console.groq.com/keys",
    },
    "OpenRouter": {
        "id": "openrouter",
        "badge": "badge-openrouter",
        "models": [
            # ── 有料モデル ──────────────────────────────────────
            "meta-llama/llama-4-maverick",
            "anthropic/claude-sonnet-4-6",
            "anthropic/claude-opus-4-6-fast",
            "openai/gpt-4o",
            "openai/o3-mini",
            "google/gemini-2.0-flash-001",
            "google/gemini-2.5-pro-preview",
            "deepseek/deepseek-chat-v3-0324",
            "mistralai/mistral-large",
            # ── 無料モデル（:free） ─────────────────────────────
            "[無料] deepseek/deepseek-v4-flash:free",
            "[無料] nvidia/nemotron-3-super-120b-a12b:free",
            "[無料] openai/gpt-oss-120b:free",
            "[無料] meta-llama/llama-3.3-70b-instruct:free",
            "[無料] google/gemma-4-31b-it:free",
            "[無料] minimax/minimax-m2.5:free",
            "[無料] qwen/qwen3-coder:free",
            "[無料] moonshotai/kimi-k2.6:free",
            "[無料] qwen/qwen3-next-80b-a3b-instruct:free",
            "[無料] openai/gpt-oss-20b:free",
            "[無料] z-ai/glm-4.5-air:free",
            "[無料] nousresearch/hermes-3-llama-3.1-405b:free",
            "[無料] openrouter/free",
        ],
        "key_placeholder": "sk-or-...",
        "key_url": "https://openrouter.ai/keys",
    },
}

# ─── セクション定義 ───────────────────────────────────────────────────────────
SECTIONS = [
    {
        "id": "basic_info",
        "title": "申請者情報（様式第15号）",
        "fields": [
            "氏名又は名称",
            "代表者の氏名",
            "住所又は主たる事務所の所在地",
            "担当者名・連絡先（電話・メール）",
            "業種（日本標準産業分類）",
            "申請日",
        ],
    },
    {
        "id": "project_overview",
        "title": "事業概要（様式第16号 第１章①）",
        "fields": [
            "事業名称",
            "対象営農類型・農作業（水田作／畑作等）",
            "対象品目名",
            "現状・課題・技術ニーズ",
            "開発供給事業の概要",
            "事業期間（実施開始年月～目標年度）",
        ],
    },
    {
        "id": "development_plan",
        "title": "開発段階の取組（様式第16号 第１章（ⅲ））",
        "fields": [
            "開発する技術の内容",
            "生産性向上目標（数値）",
            "目標数値の計算方法・根拠",
            "開発スケジュール（年度・四半期別）",
            "開発体制・人員計画",
        ],
    },
    {
        "id": "supply_plan",
        "title": "供給段階の取組（様式第16号 第１章（ⅳ））",
        "fields": [
            "供給する農業資材又はサービスの内容・供給方法",
            "供給目標数量等",
            "目標数値の根拠",
            "品質面・費用面での優位性",
            "事業の持続性（経済合理性）",
            "農業者への情報提供・継続使用支援措置",
        ],
    },
    {
        "id": "implementation",
        "title": "実施体制（様式第16号 第２章）",
        "fields": [
            "実施体制・役割分担",
            "モニタリング・評価方法",
            "リスクと対応策",
        ],
    },
    {
        "id": "budget",
        "title": "資金計画（別表４）",
        "fields": [
            "総事業費（千円）",
            "設備投資額の内訳（年度別）",
            "運転資金額の内訳（年度別）",
            "資金調達方法（補助金・借入・自己資金等）",
            "収支見通し",
        ],
    },
]

SYSTEM_PROMPT = """
【制度の本質・審査の根幹】
この制度は「人口減少下においても安定した食料供給体制を維持する」ための施策である。
農業従事者が減少しても生産水準を維持・向上できるよう、少ない人数・少ない労働力で同等以上の生産を実現するスマート農業技術の開発・供給を認定するものである。

【認定される技術の本質】
認定に値する技術とは次のいずれかを実現するものである:
(A) 省力化: 3人必要だった作業を1人でできる、10時間かかっていた作業が3時間になる → 人手不足下でも生産を継続できる
(B) 収量・可販収量の向上: 同じ労働力でより多くの農産物を生産・出荷できる → 食料供給量の確保
(C) 品質向上: より多くの国民・消費者に裨益する場合に限り有効（例：規格外品の削減で廃棄ロスを減らし可販収量を増やす等）

【認定されない・評価が低い技術の典型例（必ず指摘すること）】
× 軽労化のみ: 「体への負担が楽になる」「疲れにくくなる」だけでは不可。作業時間・人数が変わらなければ省力化にならない。人口減少対策にならない
× 付加価値額向上のみ: 「品質が上がったので倍の値段で売れる」「農家の収入が増える」だけでは不可。生産量・供給量が変わらなければ食料安全保障に貢献しない。国民が農産物に支払う金額が不当に高くなるだけで、国費を投じる公益性がない
× 生産量も労働時間も変わらず農業者の収益だけが増える技術: 国民・消費者への裨益がなく、国費をかけて認定する公益的根拠がない
× 既存技術の組み合わせで実質的な新規開発がない技術

【評価の問い直し（申請者が以下を主張したら必ず確認すること）】
・「軽労化になる」→「それによって何人が何人になりますか？何時間が何時間になりますか？」と省力化の数字を求める
・「品質が向上する」→「それによって生産量・可販収量はどう変わりますか？より多くの消費者に届きますか？」と問い直す
・「付加価値額が上がる」→「それは生産量の増加によるものですか？それとも単価上昇のみですか？単価上昇のみなら公益的根拠を示してください」と問い直す
・「農家の収入が増える」→「それは消費者への食料供給増加を伴いますか？」と問い直す

【基本姿勢】
- 審査官の目線で、要件を満たさない回答・数字・論理には明確に指摘する（おだてない）
- 申請者が道を外れそうなときは、制度の趣旨（人口減少下での食料供給維持）に立ち返って軌道修正を促す
- 曖昧な回答はそのまま受け取らず「この数字では審査を通りません」「根拠が不明確です」と具体的に指摘する
- 指摘するときは必ず「なぜ問題か（制度上の根拠）」と「どう修正すればよいか」をセットで示す
- 本制度に適さない事業は早期に指摘し、適切な他制度を案内する

【会話の進め方（この順番を守ること）】
STEP 1: スマート農業技術該当性の確認（最初に必ず行う）
STEP 2: 事業の中身を収集する（下記の認定要件に沿って）
STEP 3: 事業内容がひと通り揃ったら申請者情報を確認する

【STEP 1: スマート農業技術該当性の確認（認定の大前提）】
法第2条第1項の定義: 農業機械等の自動走行・自動制御技術、センシング・情報処理技術、その他先端的な技術を活用して農業の生産性向上を図るための技術。
対象: 農業機械等、種苗、肥料、農薬、農業用ソフトウェアその他の農業資材も含む。
確認すること:
- 開発しようとする技術がスマート農業技術に該当するか
- 該当しない場合、生産方式革新事業活動に資する先端的な技術として整理できるか
- どちらにも該当しない場合は「本制度の対象外である可能性が高い」と明示し、適切な他制度を案内する

【STEP 2: 収集する内容と認定要件】

1. 促進目標への該当確認（認定要件1・最重要）
促進目標一覧（基本方針第2の1(2)）:
・水田作/育苗及び田植: 労働時間80%削減
・水田作/除草: 労働時間80%削減
・水田作/収穫・運搬・調製: 労働時間20%削減
・畑作/播種及び移植: 労働時間60%削減
・畑作/除草: 労働時間80%削減
・畑作/収穫・運搬・選別・調製: 労働時間20%削減
・露地野菜・花き作/除草及び防除: 労働時間80%削減
・露地野菜・花き作/収穫及び運搬: 労働時間80%削減
・露地野菜・花き作/選別・調製・出荷: 労働時間60%削減
・施設野菜・花き作/栽培管理: 労働時間60%削減または付加価値額30%向上（※後者は可販収量増加や品質向上による消費者裨益が必要。単価上昇のみは不可）
・施設野菜・花き作/収穫及び運搬: 労働時間60%削減
・施設野菜・花き作/選別・調製・出荷: 労働時間60%削減または付加価値額20%向上（※同上）
・果樹・茶作/栽培管理: 労働時間60%削減
・果樹・茶作/除草及び防除: 労働時間80%削減
・果樹・茶作/収穫及び運搬: 労働時間60%削減
・果樹・茶作/選別・調製・出荷: 労働時間60%削減または付加価値額20%向上（※同上）
・畜産・酪農/飼養管理: 労働時間60%削減
・畜産・酪農/搾乳: 労働時間60%削減
・農作業共通/センシング連動: 労働時間20%削減または付加価値額20%向上（※付加価値額向上の場合は消費者裨益の根拠が必要）
・農作業共通/自動制御・遠隔操作: 労働時間40%削減
・農作業共通/熟練補助（スマートグラス等）: 労働時間20%削減

「付加価値額向上」を目標とする場合の追加確認（必ず実施）:
- 付加価値額向上の内訳は何か（収量増加・可販収量増加・廃棄ロス削減・品質向上）
- 単価上昇のみによる付加価値額向上であれば「本制度の趣旨（食料供給の安定・国民への裨益）に合致しない」と指摘する
- 品質向上の場合、より多くの消費者に同価格帯で届くことが示せるか確認する

目標値との対比:
- 該当項目の促進目標値を示し、概ね達成できるか確認する
- 概ね達成できない場合は「認定されない可能性が高い」と明示し、理由と追加的な取組条件の記載に誘導する
- 「9割」という数字は対外的に使わず「概ね達成」という表現にとどめる

2. 開発する技術の内容とBefore/After（認定要件1）
- 具体的な技術・仕組み（スマート農業技術への該当性を含む）
- Before/Afterを人数・時間・収量の数字で示す（「楽になる」「負担が減る」は数字に落とし込む）
  例: 「3人→1人」「10時間→3時間」「可販収量50kg/10a→65kg/10a」
- 数字の根拠（出典・算出方法）を必ず確認する。出典不明は審査で指摘される
- 遠隔操作技術の場合は特に「同時並行作業が実現できるか」「必要人員が何人から何人に削減されるか」を具体的に示すよう求める。「手元で操作できて楽になる」だけでは40%削減の根拠にならないと明示する
- 本技術が十分に機能するための圃場条件・生産現場の条件

3. 生産方式革新への貢献（認定要件1）
- 農業者が導入時に変える必要がある栽培方法・圃場設計等のBefore/After
- 推奨する新たな生産方式の内容と周知方法

4. 供給計画（認定要件2）
- 供給内容・供給方法・販売体制
- 供給目標数量と積算根拠（「何となく○台」は不可）
- 農業者の投資対効果: 導入コスト・削減できる労働コスト・投資回収年数を具体的数字で
- 事業採算性（経済合理性）: 生産原価・販売価格・販売台数から損益分岐点・黒字化時期
- アフターサービス・継続使用支援の体制
- 生産方式の農業者への周知方法

5. 実施体制・スケジュール（認定要件6）
- 共同申請者ごとの役割分担（「連携」だけでは不可。誰が何をするか具体的に）
- 年度・四半期別スケジュール
- リスクと対応策
- 本邦に事業拠点を有していることの確認

6. 資金計画（認定要件7）
- 総事業費・設備投資額・運転資金（年度別）
- 資金調達方法（「未定」「検討中」は不可）
- 収支見通し（黒字化時期）

【収集するフィールド（DATAタグのセクションID）】
- basic_info      : 氏名又は名称、代表者の氏名、住所又は主たる事務所の所在地、担当者名・連絡先（電話・メール）、業種（日本標準産業分類）、申請日
- project_overview: 事業名称、対象営農類型・農作業（促進目標該当項目）、対象品目名、現状・課題・技術ニーズ、開発供給事業の概要、事業期間
- development_plan: 開発する技術の内容、スマート農業技術該当性の根拠、生産性向上目標（人数・時間・収量の数値）、Before/Afterの定量的根拠（出典付き）、生産方式革新への貢献、開発スケジュール、開発体制・人員計画
- supply_plan     : 供給内容・供給方法・販売体制、供給目標数量と積算根拠、農業者の投資対効果（導入コスト・削減効果・回収期間）、事業採算性（損益計算・黒字化時期）、アフターサービス体制、生産方式の周知方法
- implementation  : 実施体制・役割分担（共同申請者ごと）、モニタリング・評価方法、リスクと対応策
- budget          : 総事業費（千円）、設備投資額の内訳（年度別）、運転資金額の内訳（年度別）、資金調達方法（具体的な調達先・金額）、収支見通し（黒字化時期）

【応答ルール】
- 一度に1〜2点ずつ聞く。一気に質問を詰め込まない
- 「楽になる」「負担が減る」「品質が上がる」「収入が増える」という回答には必ず「それは何人が何人になりますか？何時間が何時間になりますか？消費者への食料供給はどう変わりますか？」と省力化・供給量増加の数字に落とし込む
- 制度の趣旨（人口減少下での食料供給維持）に合致しない方向性は早期に指摘し、軌道修正か他制度への案内を行う
- 情報が確定したら必ず以下のタグで出力する（複数可）：
  <DATA>{"section": "セクションID", "field": "フィールド名", "value": "内容"}</DATA>
- 収集済みの情報は会話履歴を参照し、重複質問しない
"""



INTRO = """こんにちは！
**開発供給実施計画 申請書作成アシスタント**です。

農林水産省の公式様式（別記様式第15号・第16号）に準拠したWord文書を、対話形式で作成します。
申請者情報などの事務的な項目は最後にまとめて確認しますので、まずは**事業の中身**からお聞きします。

---

**どのような農業課題を解決するための事業ですか？**
（例：水稲の除草作業を自動化するロボット除草機の開発・販売、施設野菜の環境制御AIサービスの提供 など、自由にお聞かせください）"""

# ─── セッション初期化 ─────────────────────────────────────────────────────────
def init_session():
    defaults = {
        "messages": [],
        "collected_data": {},
        "provider_name": "OpenRouter",
        "model": "meta-llama/llama-4-maverick",
        "api_keys": {},   # {provider_id: key}
        "uploaded_file_text": "",   # アップロードファイルのテキスト
        "uploaded_file_name": "",
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v

init_session()

# ─── AI呼び出し（マルチプロバイダー） ────────────────────────────────────────
def build_system_prompt() -> str:
    """SYSTEM_PROMPTに収集済みデータを動的付加して返す"""
    cd = st.session_state.collected_data
    if not any(cd.values()):
        return SYSTEM_PROMPT

    # 入力済み内容を整形
    lines = ["\n\n【現在の入力済み内容（直接編集モードまたは過去の会話で収集済み）】"]
    lines.append("以下の内容がすでに入力されています。これを踏まえた上で会話してください。")
    lines.append("不足している項目・内容が薄い項目・認定要件上問題がある記述を積極的に指摘・改善提案してください。\n")

    SECTION_LABELS = {
        "basic_info"      : "申請者情報",
        "project_overview": "事業概要",
        "development_plan": "開発段階の取組",
        "supply_plan"     : "供給段階の取組",
        "implementation"  : "実施体制",
        "budget"          : "資金計画",
    }
    for sec_id, sec_label in SECTION_LABELS.items():
        fields = cd.get(sec_id, {})
        if not fields:
            continue
        lines.append(f"■ {sec_label}")
        for field, value in fields.items():
            # 長すぎる場合は先頭300文字に抑える
            v = value[:300] + "…" if len(value) > 300 else value
            lines.append(f"  [{field}]: {v}")
        lines.append("")

    lines.append("【未入力の項目については、会話の中で順次確認・補完してください】")
    return SYSTEM_PROMPT + "\n".join(lines)


def call_ai(user_message: str) -> str:
    provider_name = st.session_state.provider_name
    prov = PROVIDERS[provider_name]
    pid  = prov["id"]
    model = st.session_state.model.replace("[無料] ", "")
    api_key = st.session_state.api_keys.get(pid, "")

    msgs = [{"role": m["role"], "content": m["content"]} for m in st.session_state.messages]
    msgs.append({"role": "user", "content": user_message})

    if pid == "claude":
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        resp = client.messages.create(
            model=model, max_tokens=2000,
            system=build_system_prompt(), messages=msgs,
        )
        return resp.content[0].text

    elif pid == "gemini":
        import google.generativeai as genai
        genai.configure(api_key=api_key)
        gmodel = genai.GenerativeModel(
            model_name=model,
            system_instruction=build_system_prompt(),
        )
        # Gemini はrole="model"を使う
        history = []
        for m in msgs[:-1]:
            role = "model" if m["role"] == "assistant" else "user"
            history.append({"role": role, "parts": [m["content"]]})
        chat = gmodel.start_chat(history=history)
        response = chat.send_message(msgs[-1]["content"])
        return response.text

    elif pid == "openai":
        from openai import OpenAI
        client = OpenAI(api_key=api_key)
        openai_msgs = [{"role": "system", "content": build_system_prompt()}] + msgs
        resp = client.chat.completions.create(
            model=model, max_tokens=2000, messages=openai_msgs,
        )
        return resp.choices[0].message.content

    elif pid == "openrouter":
        from openai import OpenAI
        client = OpenAI(
            api_key=api_key,
            base_url="https://openrouter.ai/api/v1",
            default_headers={"X-Data-Policy": "no-training"}
        )
        openai_msgs = [{"role": "system", "content": build_system_prompt()}] + msgs
        resp = client.chat.completions.create(
            model=model, max_tokens=2000, messages=openai_msgs,
        )
        return resp.choices[0].message.content

    elif pid == "groq":
        from openai import OpenAI
        client = OpenAI(
            api_key=api_key,
            base_url="https://api.groq.com/openai/v1",
        )
        openai_msgs = [{"role": "system", "content": build_system_prompt()}] + msgs
        resp = client.chat.completions.create(
            model=model, max_tokens=2000, messages=openai_msgs,
        )
        return resp.choices[0].message.content

    else:
        raise ValueError(f"Unknown provider: {pid}")

# ─── ユーティリティ ───────────────────────────────────────────────────────────
def extract_data_tags(text):
    results = []
    for m in re.finditer(r"<DATA>(.*?)</DATA>", text, re.DOTALL):
        try: results.append(json.loads(m.group(1).strip()))
        except: pass
    return results

def update_collected_data(items):
    for item in items:
        sec, field, value = item.get("section",""), item.get("field",""), item.get("value","")
        if sec and field and value:
            st.session_state.collected_data.setdefault(sec, {})[field] = value

def count_filled():
    total  = sum(len(s["fields"]) for s in SECTIONS)
    filled = sum(len(v) for v in st.session_state.collected_data.values())
    return filled, total


def extract_text_from_file(uploaded_file) -> str:
    """アップロードファイルからテキストを抽出する"""
    name = uploaded_file.name.lower()
    raw = uploaded_file.read()

    if name.endswith(".docx"):
        import io
        from docx import Document
        doc = Document(io.BytesIO(raw))
        lines = []
        for para in doc.paragraphs:
            if para.text.strip():
                lines.append(para.text)
        # テーブルも抽出
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append("　".join(cells))
        return "\n".join(lines)

    elif name.endswith((".txt", ".md")):
        for enc in ("utf-8", "utf-8-sig", "cp932", "shift_jis"):
            try:
                return raw.decode(enc)
            except Exception:
                continue
        return raw.decode("utf-8", errors="replace")

    else:
        return ""


def build_file_context(file_text: str, filename: str) -> str:
    """アップロードファイルの内容をAIへの文脈として整形する"""
    # 長すぎる場合は先頭3000文字に制限
    truncated = file_text[:3000] + "\n\n（※長すぎるため以降省略）" if len(file_text) > 3000 else file_text
    return (
        f"\n\n【アップロードされたファイル: {filename}】\n"
        "以下はユーザーがアップロードしたファイルの内容です。"
        "このファイルの内容を踏まえて、認定要件の観点からレビュー・壁打ちを行い、"
        "確定した情報はDATAタグで出力してください。\n\n"
        f"{truncated}"
    )

def generate_docx():
    payload = {"created_at": datetime.now().strftime("%Y年%m月%d日"), "sections": st.session_state.collected_data}
    script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_docx.js")
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
        json_path = f.name
    out_path = json_path.replace(".json", ".docx")
    try:
        result = subprocess.run(["node", script_path, json_path, out_path],
                                capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            raise RuntimeError(result.stderr or result.stdout)
        with open(out_path, "rb") as f:
            return f.read()
    finally:
        for p in [json_path, out_path]:
            if os.path.exists(p): os.unlink(p)

# ─── ヘッダー ─────────────────────────────────────────────────────────────────
filled, total = count_filled()
pct = int(filled / total * 100) if total else 0
prov_cur = PROVIDERS[st.session_state.provider_name]
badge_cls = prov_cur["badge"]

st.markdown(f"""
<div class="main-header">
  <h1>📋 開発供給実施計画 申請書作成アシスタント
    <span class="provider-badge {badge_cls}">{st.session_state.provider_name.split(" ")[0]}</span>
  </h1>
  <p>農林水産省 別記様式第15号・第16号 対応 ／ Word形式で出力</p>
  <div class="progress-bar-wrap"><div class="progress-bar-fill" style="width:{pct}%"></div></div>
  <p style="font-size:0.78rem;opacity:0.8;margin-top:0.3rem">入力進捗：{filled} / {total} 項目（{pct}%）</p>
</div>
""", unsafe_allow_html=True)

col_chat, col_preview = st.columns([3, 2], gap="large")

# ─── チャット列 ───────────────────────────────────────────────────────────────
with col_chat:

    # ── AIプロバイダー設定パネル ──────────────────────────────────────────────
    with st.expander("⚙️ AIプロバイダー設定", expanded=not st.session_state.api_keys):
        selected_provider = st.selectbox(
            "使用するAI",
            options=list(PROVIDERS.keys()),
            index=list(PROVIDERS.keys()).index(st.session_state.provider_name),
            key="provider_select",
        )
        prov_info = PROVIDERS[selected_provider]
        pid = prov_info["id"]

        selected_model = st.selectbox(
            "モデル",
            options=prov_info["models"],
            index=0,
            key=f"model_select_{pid}",
        )

        current_key = st.session_state.api_keys.get(pid, "")
        new_key = st.text_input(
            f"APIキー（[取得はこちら]({prov_info['key_url']})）",
            value=current_key,
            type="password",
            placeholder=prov_info["key_placeholder"],
            key=f"key_input_{pid}",
        )

        if st.button("このプロバイダーで使用する", type="primary", use_container_width=True):
            if new_key:
                st.session_state.api_keys[pid] = new_key
            if st.session_state.provider_name != selected_provider or st.session_state.model != selected_model:
                # プロバイダー変更時はチャット履歴をリセット
                st.session_state.messages = []
            st.session_state.provider_name = selected_provider
            st.session_state.model = selected_model
            st.success(f"✓ {selected_provider} / {selected_model} に切り替えました")
            st.rerun()

        # 設定済みプロバイダー一覧
        set_provs = [name for name, p in PROVIDERS.items() if st.session_state.api_keys.get(p["id"])]
        if set_provs:
            st.caption("設定済み: " + " / ".join(set_provs))

    # APIキー未設定チェック
    cur_pid = PROVIDERS[st.session_state.provider_name]["id"]
    if not st.session_state.api_keys.get(cur_pid):
        st.warning("⬆️ 上の設定パネルでAPIキーを入力してください")
        st.stop()

    st.subheader("💬 AIアシスタント")

    # ── ファイルアップロード ───────────────────────────────────────────────
    with st.expander("📂 ファイルをアップロード（Word / テキスト）", expanded=False):
        st.caption("既存の資料・下書きをアップロードすると、AIがレビュー・壁打ち・転記を行います。")
        uploaded = st.file_uploader(
            "ファイルを選択",
            type=["docx", "txt", "md"],
            label_visibility="collapsed",
            key="file_uploader",
        )
        col_mode1, col_mode2 = st.columns(2)
        with col_mode1:
            do_review  = st.button("📋 レビュー・壁打ち", use_container_width=True,
                                   help="AIがファイルを読んで認定要件の観点から指摘・改善提案します")
        with col_mode2:
            do_transfer = st.button("📥 フォームに自動転記", use_container_width=True,
                                    help="AIがファイルから情報を抽出してフォームに転記します")

        if uploaded and (do_review or do_transfer):
            with st.spinner("ファイルを読み込み中..."):
                file_text = extract_text_from_file(uploaded)
            if not file_text.strip():
                st.error("テキストを抽出できませんでした。ファイルを確認してください。")
            else:
                st.session_state.uploaded_file_text = file_text
                st.session_state.uploaded_file_name = uploaded.name
                file_ctx = build_file_context(file_text, uploaded.name)

                if do_review:
                    prompt = (
                        f"以下のファイル（{uploaded.name}）をアップロードしました。"
                        "認定要件（スマート農業技術該当性・促進目標への対応・省力化の数値根拠・"
                        "供給計画・事業採算性など）の観点から、問題点・不足点・改善すべき点を"
                        "具体的に指摘してください。" + file_ctx
                    )
                    mode_label = "レビュー・壁打ち"
                else:
                    prompt = (
                        f"以下のファイル（{uploaded.name}）をアップロードしました。"
                        "ファイルの内容から申請書の各項目に対応する情報を抽出し、"
                        "DATAタグを使ってフォームに転記してください。"
                        "転記後、不足している項目や補強が必要な箇所を教えてください。" + file_ctx
                    )
                    mode_label = "自動転記"

                st.session_state.messages.append({"role": "user", "content": f"[{mode_label}] {uploaded.name} をアップロードしました。"})
                with st.spinner(f"AIが{mode_label}しています..."):
                    try:
                        reply = call_ai(prompt)
                    except Exception as e:
                        reply = f"⚠️ エラー: {e}"
                st.session_state.messages.append({"role": "assistant", "content": reply})
                extracted = extract_data_tags(reply)
                if extracted:
                    update_collected_data(extracted)
                st.rerun()

    if not st.session_state.messages:
        st.session_state.messages.append({"role": "assistant", "content": INTRO})

    chat_box = st.container(height=460)
    with chat_box:
        for msg in st.session_state.messages:
            if msg["role"] == "user":
                st.markdown(f'<div class="chat-user">🧑 {msg["content"]}</div>', unsafe_allow_html=True)
            else:
                display = re.sub(r"<DATA>.*?</DATA>", "", msg["content"], flags=re.DOTALL).strip()
                st.markdown(f'<div class="chat-ai">🤖 {display}</div>', unsafe_allow_html=True)

    with st.form("chat_form", clear_on_submit=True):
        user_input = st.text_area("", height=80, placeholder="回答を入力してください...", label_visibility="collapsed")
        c1, c2 = st.columns([3, 1])
        with c1: submitted = st.form_submit_button("送信 ➤", type="primary", use_container_width=True)
        with c2: reset     = st.form_submit_button("リセット", use_container_width=True)

    if reset:
        st.session_state.messages = []
        st.session_state.collected_data = {}
        st.rerun()

    if submitted and user_input.strip():
        st.session_state.messages.append({"role": "user", "content": user_input})
        with st.spinner(f"{st.session_state.provider_name.split(' ')[0]} が考えています..."):
            try:    reply = call_ai(user_input)
            except Exception as e: reply = f"⚠️ エラー: {e}"
        st.session_state.messages.append({"role": "assistant", "content": reply})
        extracted = extract_data_tags(reply)
        if extracted: update_collected_data(extracted)
        st.rerun()

# ─── プレビュー列 ─────────────────────────────────────────────────────────────
with col_preview:
    st.subheader("📄 申請書プレビュー")

    cd = st.session_state.collected_data
    for sec in SECTIONS:
        sec_data = cd.get(sec["id"], {})
        fc = len(sec_data); tc = len(sec["fields"])
        badge = f'<span class="cbadge">✓ 完了</span>' if fc==tc else f'<span class="ibadge">{fc}/{tc}</span>'
        fields_html = "".join(
            f"<div style='font-size:0.8rem;color:#555;margin:2px 0'>"
            f"<b>{f}</b>：{sec_data.get(f, '<span style=\"color:#bbb\">未入力</span>')}</div>"
            for f in sec["fields"]
        )
        st.markdown(f'<div class="section-preview"><div class="section-title">{sec["title"]} {badge}</div>{fields_html}</div>',
                    unsafe_allow_html=True)

    st.markdown("---")

    if filled > 0:
        if st.button("📥 Word申請書を生成", type="primary", use_container_width=True):
            with st.spinner("Word文書を生成中..."):
                try:
                    docx_bytes = generate_docx()
                    fname = f"開発供給実施計画申請書_{datetime.now().strftime('%Y%m%d')}.docx"
                    st.download_button("⬇️ 申請書.docx をダウンロード", data=docx_bytes,
                                       file_name=fname,
                                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                       use_container_width=True)
                except Exception as e:
                    st.error(f"Word生成エラー: {e}")
    else:
        st.info("対話で情報を入力するとWordダウンロードが可能になります")

    filled_count = sum(len(v) for v in cd.values())
    edit_label = f"✏️ 直接入力モード（{filled_count}項目入力済み）" if filled_count else "✏️ 直接入力モード"
    with st.expander(edit_label):
        st.caption("ここで入力した内容はチャットのAIが自動的に参照します。入力後にチャットで「レビューして」と話しかけてください。")
        sel = st.selectbox("セクション", [s["title"] for s in SECTIONS], label_visibility="collapsed")
        sec_obj = next(s for s in SECTIONS if s["title"] == sel)
        sid = sec_obj["id"]
        cd.setdefault(sid, {})
        for field in sec_obj["fields"]:
            val     = cd[sid].get(field, "")
            new_val = st.text_area(field, value=val, height=68, key=f"d_{sid}_{field}")
            if new_val != val:
                cd[sid][field] = new_val
